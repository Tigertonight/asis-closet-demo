import io
import json
import uuid
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import product_collection as pc
from app.auth import get_admin_user


@pytest.fixture
def intake(monkeypatch, tmp_path):
    monkeypatch.setattr(pc, 'STORE', tmp_path / 'store')
    app = FastAPI(); app.include_router(pc.router)
    return app, TestClient(app)


def data(**kwargs):
    return {'brand': '测试品牌', 'product_name': '测试商品', 'category': '食品',
            'contact': '测试联系人', 'contact_method': 'test_wechat', **kwargs}


def image_bytes():
    image = Image.new('RGB', (800, 600), 'white')
    result = io.BytesIO(); image.save(result, 'PNG')
    return result.getvalue()


def submit(c, values=None, image=None, request_id=None):
    body = {'payload': json.dumps(values if values is not None else data()), 'request_id': request_id or str(uuid.uuid4())}
    return c.post('/api/v1/selfit/product-collection', data=body,
                  files={'photo': ('商品.png', image, 'image/png')} if image is not None else None)


def test_submit_persists_all_fields_and_private_original(intake):
    app, c = intake
    values = data(**{k: ('yes' if kind == 'status' else '1.25' if kind == 'money' else '0' if kind == 'integer' else 'https://example.com/item' if kind == 'url' else '填写内容') for k, _, kind, _, _ in pc.FIELDS if k not in data() and kind not in {'image','selling'}})
    values['selling_points'] = {k: label + '内容' for k, label in pc.SELLING_FIELDS}
    raw = image_bytes(); saved = submit(c, values, raw); assert saved.status_code == 200
    identifier = saved.json()['id']
    for path in ['/admin/api/product-collection', '/admin/api/product-collection/export.xlsx', f'/admin/api/product-collection/{identifier}', f'/admin/api/product-collection/{identifier}/image', f'/admin/api/product-collection/{identifier}/image?preview=true']:
        assert c.get(path).status_code == 401
        assert c.get(path, headers={'Authorization':'Bearer non-admin-token'}).status_code in (401,403)
    app.dependency_overrides[get_admin_user] = lambda: {'user_id': 'admin_console'}
    assert c.get('/admin/api/product-collection').json()['total'] == 1
    assert c.get('/admin/api/product-collection').json()['today'] == 1
    record = c.get(f'/admin/api/product-collection/{identifier}'); assert record.json()['data'] == values
    assert record.headers['cache-control'] == 'private, no-store'
    original = c.get(f'/admin/api/product-collection/{identifier}/image')
    assert original.content == raw and 'attachment;' in original.headers['content-disposition']
    preview = c.get(f'/admin/api/product-collection/{identifier}/image?preview=true')
    with Image.open(io.BytesIO(preview.content)) as img:
        assert img.format == 'JPEG' and max(img.size) == 240
    # Reading via a fresh connection returns the same durable record.
    with pc.database() as db:
        assert db.execute('SELECT count(*) FROM products').fetchone()[0] == 1


def test_idempotency_including_concurrent_retry(intake):
    app, c = intake; key = str(uuid.uuid4()); raw=image_bytes()
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses=list(pool.map(lambda _:pc.save_submission(key, pc.clean_data(json.dumps(data())), raw),range(4)))
    assert len({r['id'] for r in responses}) == 1
    assert submit(c, image=raw, request_id=key).status_code == 200
    assert submit(c, data(product_name='changed'), image=raw, request_id=key).status_code == 409
    with pc.database() as db:
        assert db.execute('SELECT count(*) FROM products').fetchone()[0] == 1
        assert db.execute('SELECT count(*) FROM product_images').fetchone()[0] == 1


@pytest.mark.parametrize('values', [data(brand=' '),data(contact_method=''),data(daily_price='-1'),data(live_price='1.001'),data(stock='1.5'),data(lottery_stock='1000000000'),data(product_url='javascript:alert(1)'),data(live_url='http://u:p@example.com'),data(listed='maybe'),data(selling_points={'unexpected':'x'}),data(brand='a\x00'),data(extra='x'),data(product_id=123),data(selling_points={'features':'x'*1001})])
def test_reject_invalid_fields(intake,values):
    _,c=intake;assert submit(c,values).status_code==422


def test_image_errors_and_no_partial_row(intake, monkeypatch):
    _, c = intake
    assert submit(c, image=b'not a photo').status_code == 422
    assert submit(c, image=b'').status_code == 422
    monkeypatch.setattr(pc, 'MAX_IMAGE_BYTES', 10)
    assert submit(c, image=image_bytes()).status_code == 413
    with pc.database() as db:
        assert db.execute('SELECT count(*) FROM products').fetchone()[0] == 0


def test_pagination_filter_and_export_with_safe_text(intake, monkeypatch):
    app,c=intake; app.dependency_overrides[get_admin_user]=lambda: {'user_id':'admin_console'}
    submit(c, data(brand='=HYPERLINK("bad")', product_id='00012345678901234567890', daily_price='23.50', live_price='0', stock='12', selling_points={'features':'第一行\n第二行'}), image=image_bytes())
    submit(c, data(brand='其他品牌'));submit(c, data(brand='100%纯'))
    first=c.get('/admin/api/product-collection?limit=2').json();assert first['total']==3 and len(first['items'])==2
    submit(c, data(brand='新到达'))
    second=c.get(f'/admin/api/product-collection?limit=2&before_id={first["next_cursor"]}').json()
    assert {x['id'] for x in first['items']}.isdisjoint({x['id'] for x in second['items']})
    assert len(second['items'])==1
    assert c.get('/admin/api/product-collection',params={'q':'%'}).json()['filtered']==1
    exported=c.get('/admin/api/product-collection/export.xlsx',params={'q':'HYPERLINK'})
    assert exported.status_code==200
    workbook=load_workbook(io.BytesIO(exported.content));ws=workbook.active
    assert list(next(ws.values))[:22]==[f[1] for f in pc.FIELDS]
    assert ws.max_row==2 and ws.max_column==24
    assert ws['A2'].value=='=HYPERLINK("bad")' and ws['A2'].data_type=='s'
    assert ws['E2'].value=='00012345678901234567890' and ws['E2'].data_type=='s'
    assert ws['K2'].value==23.5 and ws['L2'].value==0 and ws['O2'].value==12
    assert '第一行\n第二行' in ws['Q2'].value
    assert len(ws._images)==1
    assert ws.freeze_panes=='C2' and ws.auto_filter.ref=='A1:X2'
    assert c.get('/admin/api/product-collection/export.xlsx',params={'q':'not found'}).status_code==200
    monkeypatch.setattr(pc,'MAX_EXPORT_ROWS',2)
    assert c.get('/admin/api/product-collection/export.xlsx').status_code==422


def test_routes_schema_and_missing_record(intake):
    app,c=intake
    assert c.get('/selfit/product-collection').status_code==200
    schema=c.get('/api/v1/selfit/product-collection/schema').json()
    assert len(schema['fields'])==22 and len(schema['selling_fields'])==10
    assert '首饰' in schema['dimensions_help']
    app.dependency_overrides[get_admin_user]=lambda: {'user_id':'admin_console'}
    assert c.get('/admin/api/product-collection/99999').status_code==404
    assert c.get('/admin/api/product-collection/99999/image').status_code==404
