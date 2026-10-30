import json
import io
import uuid

import pytest
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import product_collection as pc, product_collection_auth as auth
from app import auth as main_auth
from app import ops

PASSWORD = 'collection-test-password-123'
BASE = '/product-collection-admin/api'


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(pc, 'STORE', tmp_path / 'products')
    monkeypatch.setattr(auth, 'AUTH_DB', tmp_path / 'collection-auth.sqlite3')
    monkeypatch.setattr(main_auth, 'AUTH_STORE_PATH', tmp_path / 'main-auth.json')
    monkeypatch.setenv('SELFIT_AUTH_SECRET', 'test-secret')
    auth.configure_password(PASSWORD)
    app = FastAPI(); app.include_router(pc.router)
    @app.get('/admin-scope')
    def main_scope(user=Depends(main_auth.get_admin_user)):
        return {'ok': True}
    return TestClient(app, base_url='https://testserver')


def login(client, password=PASSWORD):
    return client.post(BASE + '/login', json={'password': password})


def valid_data(**kwargs):
    values = {k: '不适用' for k, _, kind, _, _ in pc.FIELDS if kind not in {'selling', 'image'}}
    return {**values, 'selling_points':{k:'不适用' for k,_ in pc.SELLING_FIELDS}, 'image_not_applicable':True, **kwargs}


def test_independent_login_existing_rows_and_excel(client):
    saved = client.post('/api/v1/selfit/product-collection', data={'request_id':'11111111-1111-4111-8111-111111111111','payload':json.dumps(valid_data(brand='existing', product_name='商品'))})
    assert saved.status_code == 200
    assert client.get(BASE + '/products').status_code == 401
    assert login(client, 'wrong').status_code == 401
    response = login(client)
    assert response.status_code == 200
    cookie = response.headers['set-cookie'].lower()
    assert 'httponly' in cookie and 'secure' in cookie and 'samesite=strict' in cookie and 'path=/product-collection-admin' in cookie
    assert 'access_token' not in response.json()
    assert client.get(BASE + '/session').json()['role'] == 'product_collection_admin'
    assert client.get(BASE + '/products').json()['total'] == 1
    assert client.get(BASE + '/products/1').json()['data']['brand'] == 'existing'
    export = client.get(BASE + '/products/export.xlsx')
    assert export.status_code == 200 and export.content.startswith(b'PK')
    assert export.headers['cache-control'] == 'private, no-store'
    assert client.get('/admin/api/product-collection').status_code == 404


def test_both_admin_domains_are_isolated(client):
    root_token = main_auth.issue_admin_session('127.0.0.1')
    root_headers = {'Authorization': 'Bearer ' + root_token}
    assert client.get('/admin-scope', headers=root_headers).status_code == 200
    client.cookies.set(main_auth.ADMIN_COOKIE_NAME, root_token)
    assert client.get(BASE + '/products', headers=root_headers).status_code == 401
    client.cookies.clear()
    assert login(client).status_code == 200
    token = client.cookies.get(auth.COOKIE)
    assert client.get('/admin-scope').status_code == 401
    assert client.get('/admin-scope', headers={'Authorization':'Bearer '+token}).status_code == 401
    assert client.get(BASE + '/products').status_code == 200
    with auth.database() as db:
        digest = db.execute('SELECT digest FROM password').fetchone()[0]
        session_hash = db.execute('SELECT token_hash FROM sessions').fetchone()[0]
    assert PASSWORD not in digest and token != session_hash


def test_logout_rotation_and_expiry_revoke_access(client):
    login(client); token = client.cookies.get(auth.COOKIE)
    assert client.post(BASE + '/logout').status_code == 200
    client.cookies.set(auth.COOKIE, token, path='/product-collection-admin')
    assert client.get(BASE + '/products').status_code == 401
    client.cookies.clear(); login(client)
    auth.configure_password('another-collection-password-456')
    assert client.get(BASE + '/session').status_code == 401
    assert login(client).status_code == 401
    assert login(client, 'another-collection-password-456').status_code == 200
    with auth.database() as db:
        db.execute('UPDATE sessions SET expires=0')
    assert client.get(BASE + '/products/export.xlsx').status_code == 401


def test_unconfigured_and_cross_origin_fail_closed(client):
    assert client.post(BASE + '/login', json={'password':PASSWORD}, headers={'Origin':'https://other.example'}).status_code == 403
    assert client.post(BASE + '/login', json={'password':PASSWORD}, headers={'Sec-Fetch-Site':'cross-site'}).status_code == 403
    assert login(client).status_code == 200
    assert client.post(BASE + '/logout', headers={'Origin':'https://other.example'}).status_code == 403
    assert client.get(BASE + '/session').status_code == 200
    with auth.database() as db:
        db.execute('DELETE FROM password')
    assert login(client).status_code == 503


def test_separate_login_rate_rule_and_brand_free_page(client):
    rule=ops._matching_rule(BASE + '/login','POST')
    assert rule.name != ops._matching_rule('/admin/api/login','POST').name
    assert rule.max_requests == 20
    html=client.get('/product-collection-admin').text
    assert '小红书特产之夜' in html and '商品信息管理' in html
    assert '<title>小红书特产之夜 · 商品信息管理</title>' in html
    assert 'selfit' not in html.replace('/selfit/product-collection','')


def test_login_requests_are_rate_limited(monkeypatch, tmp_path):
    monkeypatch.setattr(auth, 'AUTH_DB', tmp_path / 'rate-auth.sqlite3')
    monkeypatch.setattr(ops, '_limiter', ops.InMemoryRateLimiter())
    monkeypatch.setenv('SELFIT_DISABLE_RATE_LIMIT', '0')
    app=FastAPI();app.include_router(pc.router)
    app.middleware('http')(ops.request_guard_middleware)
    c=TestClient(app)
    for _ in range(20):
        assert login(c).status_code == 503  # Missing configuration still consumes an attempt.
    response=login(c)
    assert response.status_code == 429 and int(response.headers['retry-after']) > 0


def test_delete_requires_independent_admin_and_excludes_all_reads(client):
    image = io.BytesIO(); Image.new('RGB', (40, 40), 'white').save(image, 'PNG')
    key = str(uuid.uuid4())
    fields = {'request_id': key, 'payload': json.dumps(valid_data(brand='删除验收', product_name='待删除商品', image_not_applicable=False))}
    files = {'photo': ('product.png', image.getvalue(), 'image/png')}
    identifier = client.post('/api/v1/selfit/product-collection', data=fields, files=files).json()['id']
    retained_fields = {**fields, 'request_id': str(uuid.uuid4()), 'payload': json.dumps({**json.loads(fields['payload']), 'product_name': '保留商品'})}
    retained = client.post('/api/v1/selfit/product-collection', data=retained_fields, files=files).json()['id']
    path = BASE + f'/products/{identifier}'
    assert client.delete(path).status_code == 401
    root_token = main_auth.issue_admin_session('127.0.0.1')
    assert client.delete(path, headers={'Authorization': 'Bearer ' + root_token}).status_code == 401
    assert login(client).status_code == 200
    assert client.delete(path, headers={'Origin': 'https://other.example'}).status_code == 403
    assert client.delete(path, headers={'Sec-Fetch-Site': 'cross-site'}).status_code == 403
    assert client.get(BASE + '/products').json()['total'] == 2
    deleted = client.delete(path, headers={'Origin': 'https://testserver'})
    assert deleted.status_code == 200 and deleted.json()['id'] == identifier
    assert deleted.headers['cache-control'] == 'private, no-store'
    # Retries are idempotent and keep the original deletion time.
    with pc.database() as db:
        timestamp = db.execute('SELECT deleted_at FROM product_deletions WHERE product_id=?', (identifier,)).fetchone()[0]
    assert client.delete(path).status_code == 200
    assert client.delete(BASE + '/products/99999').status_code == 404
    for suffix in ['', '/image', '/image?preview=true']:
        assert client.get(path + suffix).status_code == 404
    listing = client.get(BASE + '/products?limit=1&q=删除验收').json()
    assert listing['total'] == listing['today'] == listing['filtered'] == 1
    assert listing['next_cursor'] is None and [r['id'] for r in listing['items']] == [retained]
    assert client.get(BASE + '/products?q=待删除商品').json()['filtered'] == 0
    assert client.get(BASE + f'/products/{retained}/image').content == image.getvalue()
    workbook = load_workbook(io.BytesIO(client.get(BASE + '/products/export.xlsx').content))
    assert workbook.active.max_row == 2 and workbook.active['B2'].value == '保留商品'
    assert len(workbook.active._images) == 1
    workbook.close()
    # A late public submission retry must not resurrect a deleted record.
    assert client.post('/api/v1/selfit/product-collection', data=fields, files=files).status_code == 409
    with pc.database() as db:
        assert db.execute('SELECT count(*) FROM products').fetchone()[0] == 2
        assert db.execute('SELECT count(*) FROM product_images').fetchone()[0] == 2
        assert db.execute('SELECT deleted_at FROM product_deletions WHERE product_id=?', (identifier,)).fetchone()[0] == timestamp
    client.post(BASE + '/logout')
    assert client.delete(BASE + f'/products/{retained}').status_code == 401


def test_delete_table_upgrade_keeps_existing_records(client):
    key = str(uuid.uuid4())
    fields = {'request_id':key, 'payload':json.dumps(valid_data(brand='旧记录',product_name='旧商品'))}
    identifier = client.post('/api/v1/selfit/product-collection', data=fields).json()['id']
    with pc.database() as db:
        db.execute('DROP TABLE product_deletions')  # Schema before deletion support.
    assert login(client).status_code == 200
    assert client.get(BASE + '/products').json()['total'] == 1
    assert client.delete(BASE + f'/products/{identifier}').status_code == 200
    assert client.get(BASE + '/products').json()['total'] == 0
