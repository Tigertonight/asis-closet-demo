"""Standalone merchant intake. Each submission is an atomic SQLite transaction.

Contact details and source images are private: only administrators may list,
read images or export. No submission contents are sent to analytics/logging.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlsplit

import anyio
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from PIL import Image, ImageOps, UnidentifiedImageError
from app.auth import get_admin_user
from app.photo_work import photo_admission, run_photo_work

ROOT = Path(__file__).resolve().parents[1]
STORE = ROOT / 'outputs' / 'product_collection'
router = APIRouter(tags=['product-collection'])
PRIVATE = {'Cache-Control': 'private, no-store', 'X-Content-Type-Options': 'nosniff'}
MAX_IMAGE_BYTES = 20 * 1024 * 1024
MAX_EXPORT_ROWS = 5000
_export_slots = anyio.CapacityLimiter(1)
SELLING_FIELDS = [
    ('introduction', '一句话商品介绍'), ('features', '卖点信息'),
    ('differentiation', '独家差异化卖点'), ('shelf_life', '保质期'),
    ('production_date', '生产日期'), ('celebrity', '明星/博主同款（如有请标注）'),
    ('pregnancy', '孕妇是否可用'), ('sensitive_skin', '敏肌是否可用'),
    ('color', '颜色'), ('flavor', '口味（食品/饮品）'),
]
# Keep the supplied sheet's exact 22-column order for Excel and admin details.
FIELDS = [
    ('brand', '品牌名称', 'text', 1, True),
    ('product_name', '商品名称', 'text', 1, True),
    ('category', '品类', 'text', 1, True),
    ('product_url', '商品链接', 'url', 1, False),
    ('product_id', '商品ID', 'text', 1, False),
    ('contact', '商家联系人', 'text', 1, True),
    ('contact_method', '联系电话/微信', 'text', 1, True),
    ('operator', '对应运营', 'text', 1, False),
    ('live_url', '最终直播链接', 'url', 2, False),
    ('sku', '商品规格/SKU', 'textarea', 2, False),
    ('daily_price', '日常售价（元）', 'money', 2, False),
    ('live_price', '直播到手价（元）', 'money', 2, False),
    ('promotion', '薯店长直播优惠机制', 'textarea', 2, False),
    ('gifts', '赠品/赠品数量', 'textarea', 2, False),
    ('stock', '直播可售数量', 'integer', 2, False),
    ('product_image', '产品白底图', 'image', 3, False),
    ('selling_points', '商品卖点', 'selling', 3, False),
    ('lottery_stock', '直播提供抽奖数量', 'integer', 2, False),
    ('cautions', '商品禁忌/注意事项', 'textarea', 3, False),
    ('listed', '是否完成上架', 'status', 4, False),
    ('price_verified', '是否完成价格校验', 'status', 4, False),
    ('dimensions', '尺寸', 'textarea', 3, False),
]
DIMENSIONS_HELP = '服饰：提供具体衣长\n首饰：口径 + 外包装\n家居：长 × 宽 × 高（mm）+ 外包装尺寸（如有）\n食品：外包装长 × 宽 × 高（mm）'


@contextmanager
def database():
    STORE.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(STORE / 'products.sqlite3', timeout=10)
    db.row_factory = sqlite3.Row
    try:
        db.execute('PRAGMA journal_mode=WAL')
        db.execute('PRAGMA foreign_keys=ON')
        db.executescript('''
            CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                request_id TEXT UNIQUE NOT NULL, digest TEXT NOT NULL,
                created_at TEXT NOT NULL, brand TEXT NOT NULL,
                product_name TEXT NOT NULL, category TEXT NOT NULL, data TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS products_created ON products(created_at);
            CREATE TABLE IF NOT EXISTS product_images (
                product_id INTEGER PRIMARY KEY REFERENCES products(id),
                original BLOB NOT NULL, preview BLOB NOT NULL,
                mime TEXT NOT NULL, extension TEXT NOT NULL
            );
        ''')
        with db:
            yield db
    finally:
        db.close()


def clean_data(raw: str) -> dict:
    try:
        data = json.loads(raw)
    except ValueError:
        raise HTTPException(422, '无法读取填写内容，请刷新后重试。')
    allowed = {f[0] for f in FIELDS if f[2] != 'image'}
    if not isinstance(data, dict) or set(data) - allowed:
        raise HTTPException(422, '表单内容不正确，请刷新后重试。')
    result = {}
    for key, label, kind, _, required in FIELDS:
        if kind == 'image':
            continue
        value = data.get(key, {} if kind == 'selling' else '')
        if kind == 'selling':
            if not isinstance(value, dict) or set(value) - {k for k, _ in SELLING_FIELDS}:
                raise HTTPException(422, '商品卖点内容不正确。')
            result[key] = {k: clean_text(value.get(k, ''), label_text, 1000) for k, label_text in SELLING_FIELDS}
            continue
        value = clean_text(value, label, 2000 if kind in {'url', 'textarea'} else 200)
        if required and not value:
            raise HTTPException(422, f'请填写{label}。')
        if value and kind == 'url':
            try:
                parsed = urlsplit(value)
                valid = parsed.scheme in {'http', 'https'} and parsed.hostname and not parsed.username and not parsed.password and not re.search(r'\s', value)
            except ValueError:
                valid = False
            if not valid:
                raise HTTPException(422, f'{label}请填写完整的 http 或 https 地址。')
        if value and kind == 'money' and not re.fullmatch(r'\d{1,9}(\.\d{1,2})?', value, re.ASCII):
            raise HTTPException(422, f'{label}请填写非负金额，最多两位小数。')
        if value and kind == 'integer' and not re.fullmatch(r'\d{1,9}', value, re.ASCII):
            raise HTTPException(422, f'{label}请填写非负整数。')
        if kind == 'status' and value not in {'', 'yes', 'no'}:
            raise HTTPException(422, f'请重新选择{label}。')
        result[key] = value
    return result


def clean_text(value, label, limit):
    if not isinstance(value, str) or len(value) > limit or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', value):
        raise HTTPException(422, f'{label}内容无效或超过 {limit} 字。')
    return value.strip()


def prepare_image(raw):
    if not raw:
        return None
    try:
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in {'PNG', 'JPEG', 'WEBP'} or image.width * image.height > 40_000_000 or getattr(image, 'n_frames', 1) != 1:
                raise ValueError()
            fmt = image.format
            image.verify()
        with Image.open(io.BytesIO(raw)) as image:
            # Small derivative for Excel/admin, source bytes remain untouched.
            image.thumbnail((240, 240))
            preview = ImageOps.exif_transpose(image).convert('RGBA')
            background = Image.new('RGB', preview.size, 'white')
            background.paste(preview, mask=preview.getchannel('A'))
            buf = io.BytesIO()
            background.save(buf, 'JPEG', quality=85)
        return (raw, buf.getvalue(), {'PNG': 'image/png', 'JPEG': 'image/jpeg', 'WEBP': 'image/webp'}[fmt], {'PNG': 'png', 'JPEG': 'jpg', 'WEBP': 'webp'}[fmt])
    except (OSError, ValueError, UnidentifiedImageError, Image.DecompressionBombError):
        raise HTTPException(422, '请上传有效的 JPG、PNG 或 WebP 静态图片，尺寸不超过 4000 万像素。')


def save_submission(request_id, data, raw):
    encoded = json.dumps(data, ensure_ascii=False, sort_keys=True)
    digest = hashlib.sha256(encoded.encode() + b'\0' + raw).hexdigest()
    with database() as db:
        existing = db.execute('SELECT id,digest FROM products WHERE request_id=?', (request_id,)).fetchone()
        if existing:
            if existing['digest'] != digest:
                raise HTTPException(409, '上一次提交已收到，请刷新页面后登记新的商品。')
            return {'id': existing['id'], 'message': '提交成功，我们已收到商品信息。'}
    image = prepare_image(raw)
    with database() as db:
        # The unique key makes concurrent retries atomic as well.
        db.execute('INSERT OR IGNORE INTO products(request_id,digest,created_at,brand,product_name,category,data) VALUES(?,?,?,?,?,?,?)',
                   (request_id, digest, datetime.now(timezone.utc).isoformat(), data['brand'], data['product_name'], data['category'], encoded))
        inserted = db.execute('SELECT changes()').fetchone()[0] == 1
        row = db.execute('SELECT id,digest FROM products WHERE request_id=?', (request_id,)).fetchone()
        if row['digest'] != digest:
            raise HTTPException(409, '提交内容已变化，请刷新后重试。')
        if inserted and image:
            db.execute('INSERT INTO product_images VALUES(?,?,?,?,?)', (row['id'], *image))
        return {'id': row['id'], 'message': '提交成功，我们已收到商品信息。'}


@router.get('/selfit/product-collection', include_in_schema=False)
@router.get('/selfit/product-collection/', include_in_schema=False)
def collection_page():
    return FileResponse(ROOT / 'app/static/product-collection/index.html', headers=PRIVATE)


@router.get('/api/v1/selfit/product-collection/schema')
def collection_schema():
    return {'title': '小红书特产之夜', 'fields': [dict(key=k, label=l, type=t, group=g, required=r) for k, l, t, g, r in FIELDS],
            'selling_fields': [dict(key=k, label=l) for k, l in SELLING_FIELDS], 'dimensions_help': DIMENSIONS_HELP}


@router.post('/api/v1/selfit/product-collection')
async def submit_product(payload: str = Form(..., max_length=32000), request_id: str = Form(..., max_length=40),
                         photo: UploadFile | None = File(None)):
    try:
        request_id = str(uuid.UUID(request_id))
    except ValueError:
        raise HTTPException(422, '提交标识无效，请刷新后重试。')
    data = clean_data(payload)
    async with photo_admission():
        raw = await photo.read(MAX_IMAGE_BYTES + 1) if photo and photo.filename else b''
        if len(raw) > MAX_IMAGE_BYTES:
            raise HTTPException(413, '图片不能超过 20MB，请选择较小的图片后重试。')
        if photo and photo.filename and not raw:
            raise HTTPException(422, '图片文件为空，请重新选择。')
        saved = await run_photo_work(save_submission, request_id, data, raw)
    return JSONResponse(saved, headers=PRIVATE)


def where_query(q):
    if not q:
        return '', ()
    escaped = q.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
    return "WHERE (brand LIKE ? ESCAPE '\\' OR product_name LIKE ? ESCAPE '\\' OR category LIKE ? ESCAPE '\\')", (f'%{escaped}%',) * 3


@router.get('/admin/api/product-collection')
def list_products(q: str = Query('', max_length=100), before_id: int | None = Query(None, ge=1),
                  limit: int = Query(25, ge=1, le=100), admin=Depends(get_admin_user)):
    where, params = where_query(q.strip())
    cutoff = (datetime.now(timezone(timedelta(hours=8))).replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc).isoformat())
    with database() as db:
        total, today = db.execute('SELECT count(*),coalesce(sum(created_at>=?),0) FROM products', (cutoff,)).fetchone()
        filtered = db.execute(f'SELECT count(*) FROM products {where}', params).fetchone()[0]
        cursor_where = where + (' AND ' if where else ' WHERE ') + 'id<?' if before_id else where
        cursor_params = (*params, before_id) if before_id else params
        rows = db.execute(f'SELECT id,created_at,brand,product_name,category,EXISTS(SELECT 1 FROM product_images WHERE product_id=products.id) AS has_image FROM products {cursor_where} ORDER BY id DESC LIMIT ?', (*cursor_params, limit + 1)).fetchall()
    return JSONResponse({'total': total, 'today': today, 'filtered': filtered, 'items': [dict(r) for r in rows[:limit]],
                         'next_cursor': rows[limit - 1]['id'] if len(rows) > limit else None}, headers=PRIVATE)


def export_workbook(q):
    from tempfile import SpooledTemporaryFile
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    where, params = where_query(q.strip())
    output = SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode='w+b')
    try:
        with database() as db:
            db.execute('BEGIN')  # One read snapshot for count + export while submissions continue.
            count = db.execute(f'SELECT count(*) FROM products {where}', params).fetchone()[0]
            if count > MAX_EXPORT_ROWS:
                raise HTTPException(422, f'单次最多导出 {MAX_EXPORT_ROWS} 条，请按品牌、商品或品类筛选后导出。')
            wb = Workbook(write_only=True)
            sheet = wb.create_sheet('直播商品信息')
            sheet.freeze_panes = 'C2'
            headers = [f[1] for f in FIELDS] + ['提交时间（北京时间）', '记录编号']
            for i, header in enumerate(headers, 1):
                sheet.column_dimensions[get_column_letter(i)].width = 48 if header in {'商品卖点', '尺寸', '商品禁忌/注意事项'} else 24
            def cell(value, heading=False):
                c = WriteOnlyCell(sheet, value=value)
                # All strings are text, never Excel formulas (also preserves long IDs/leading zeros).
                if isinstance(value, str):
                    c.data_type = 's'
                c.alignment = Alignment(vertical='top', wrap_text=True)
                if heading:
                    c.font = Font(bold=True, color='FFFFFF')
                    c.fill = PatternFill('solid', fgColor='8A011B')
                return c
            sheet.append([cell(h, True) for h in headers])
            cursor = db.execute(f'SELECT id,data,created_at FROM products {where} ORDER BY id DESC', params)
            for row_num, row in enumerate(cursor, 2):
                data = json.loads(row['data'])
                image = db.execute('SELECT preview FROM product_images WHERE product_id=?', (row['id'],)).fetchone()
                values = []
                for key, _, kind, _, _ in FIELDS:
                    value = data.get(key, '')
                    if kind == 'selling':
                        value = '\n'.join(f'{label}：{value.get(k, "")}' for k, label in SELLING_FIELDS)
                    elif kind == 'status':
                        value = {'yes': '是', 'no': '否', '': '待确认'}[value]
                    elif kind == 'image':
                        value = '缩略图（高清原图请在后台下载）' if image else ''
                    elif value and kind in {'money', 'integer'}:
                        value = float(value) if kind == 'money' else int(value)
                    values.append(cell(value))
                created = datetime.fromisoformat(row['created_at']).astimezone(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')
                sheet.row_dimensions[row_num].height = 180
                sheet.append(values + [cell(created), cell(row['id'])])
                if image:
                    drawing = XLImage(io.BytesIO(image['preview']))
                    scale = min(1, 150 / max(drawing.width, drawing.height))
                    drawing.width *= scale; drawing.height *= scale
                    sheet.add_image(drawing, f'P{row_num}')
            sheet.auto_filter.ref = f'A1:X{count + 1}'
            wb.save(output)
        output.seek(0)
        return output
    except BaseException:
        output.close()
        raise


@router.get('/admin/api/product-collection/export.xlsx')
async def export_products(q: str = Query('', max_length=100), admin=Depends(get_admin_user)):
    # Bound simultaneous workbook creation independently from image decoding.
    async with _export_slots:
        result = await anyio.to_thread.run_sync(export_workbook, q)
    def chunks():
        try:
            while chunk := result.read(65536):
                yield chunk
        finally:
            result.close()
    filename = 'products-' + datetime.now(timezone(timedelta(hours=8))).strftime('%Y%m%d-%H%M%S') + '.xlsx'
    return StreamingResponse(chunks(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             headers={**PRIVATE, 'Content-Disposition': f'attachment; filename="{filename}"'})


@router.get('/admin/api/product-collection/{product_id}')
def product_detail(product_id: int, admin=Depends(get_admin_user)):
    with database() as db:
        row = db.execute('SELECT id,created_at,data FROM products WHERE id=?', (product_id,)).fetchone()
        if row is None:
            raise HTTPException(404, '未找到这条商品信息。')
        image = db.execute('SELECT length(original) AS bytes,mime FROM product_images WHERE product_id=?', (product_id,)).fetchone()
    return JSONResponse({'id': row['id'], 'created_at': row['created_at'], 'data': json.loads(row['data']), 'image': dict(image) if image else None}, headers=PRIVATE)


@router.get('/admin/api/product-collection/{product_id}/image')
def product_image(product_id: int, preview: bool = False, admin=Depends(get_admin_user)):
    with database() as db:
        row = db.execute(f'SELECT {"preview" if preview else "original"} AS content,mime,extension FROM product_images WHERE product_id=?', (product_id,)).fetchone()
    if row is None:
        raise HTTPException(404, '没有找到商品图片。')
    return Response(row['content'], media_type='image/jpeg' if preview else row['mime'], headers={**PRIVATE,
        'Content-Disposition': 'inline' if preview else f'attachment; filename="product-{product_id}.{row["extension"]}"'})
